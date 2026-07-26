#!/usr/bin/env python3
"""把框图重建为完全可编辑的 PPT 原生形状（配方来自 ppt_study 实测规范）。"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import framework_data as D

EMU_IN = 914400
FIG_H = D.H
CMB = set('̂̄̇')  # 组合帽/杠/点 → Cambria Math 更稳


def ty(y_mpl, h=0.0):
    return FIG_H - y_mpl - h


def set_run(r, text, size=9, bold=False, italic=False, color='000000',
            baseline=None):
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    latin = ('Cambria Math' if any(c in CMB for c in text)
             or any(c in '⊖ℝ' for c in text) else 'Times New Roman')
    rPr = r._r.get_or_add_rPr()
    for tag, tf in (('a:latin', latin), ('a:ea', '宋体')):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set('typeface', tf)
    if baseline is not None:
        rPr.set('baseline', str(baseline))


def add_runs(par, runs, size, color, bold=False):
    for item in runs:
        if isinstance(item, str):
            set_run(par.add_run(), item, size=size, bold=bold, color=color)
        else:
            text, st = item
            bl = -25000 if st.get('sub') else (30000 if st.get('sup') else None)
            set_run(par.add_run(), text, size=size * (0.78 if bl else 1.0),
                    bold=st.get('b', bold), italic=st.get('i', False),
                    color=color, baseline=bl)


def arrowize(shape, color, w_pt, dash=False, tail=True):
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(w_pt)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = shape.line._get_or_add_ln()
    if tail:
        ln.append(ln.makeelement(qn('a:tailEnd'),
                                 {'type': 'triangle', 'w': 'med', 'len': 'med'}))


def polyline(shapes, pts_mpl, color, w_pt, dash=False, tail=True):
    pts = [(x, ty(y)) for x, y in pts_mpl]
    b = shapes.build_freeform(Emu(int(pts[0][0] * EMU_IN)),
                              Emu(int(pts[0][1] * EMU_IN)), scale=1.0)
    b.add_line_segments([(Emu(int(x * EMU_IN)), Emu(int(y * EMU_IN)))
                         for x, y in pts[1:]], close=False)
    ff = b.convert_to_shape()
    ff.fill.background()
    ff.shadow.inherit = False
    arrowize(ff, color, w_pt, dash=dash, tail=tail)
    return ff


def textbox(shapes, x, y, runs, size, color, align='c', rot=0, bold=False,
            width=None, height=0.28):
    # 极窄锚定框 + wrap=none：文字按段落对齐从锚定框溢出绘制，
    # 保证 左/中/右 对齐分别精确锚定在 x（LibreOffice/Office 行为一致）。
    width = 0.06
    if align == 'c':
        left = x - width / 2
    elif align == 'l':
        left = x
    else:
        left = x - width
    tb = shapes.add_textbox(Inches(left), Inches(ty(y) - height / 2),
                            Inches(width), Inches(height))
    tb.text_frame.auto_size = None
    tf = tb.text_frame
    tf.word_wrap = False
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = {'c': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT,
                   'r': PP_ALIGN.RIGHT}[align]
    if isinstance(runs, str):
        runs = [runs]
    add_runs(p, runs, size, color, bold=bold)
    if rot:
        tb.rotation = -rot
        # 旋转绕中心：把中心挪到目标点
        tb.left = Inches(x - width / 2)
        tb.top = Inches(ty(y) - height / 2)
    return tb


prs = Presentation()
prs.slide_width = Inches(D.W)
prs.slide_height = Inches(FIG_H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes

# ---------- 层带容器 ----------
for x, y, w, h, title, badge, ec, fc in D.BANDS:
    c = S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(ty(y, h)),
                    Inches(w), Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string(fc)
    c.line.color.rgb = RGBColor.from_string(ec)
    c.line.width = Pt(0.8)
    c.adjustments[0] = min(0.5, 0.15 / min(w, h))
    c.shadow.inherit = False
    textbox(S, x + 0.12, y + h - 0.21, title, 10.5, ec, align='l', bold=True,
            width=3.5)
    textbox(S, x + w - 0.12, y + h - 0.21, badge, 8, ec, align='r', width=2.6)

# ---------- 物理耦合双线 ----------
for (pts,) in D.COUPLE:
    polyline(S, pts, '888888', 2.2, tail=False)

# ---------- 连线 ----------
for pts, kind, tail in D.SEGS:
    if kind == 'rt':
        polyline(S, pts, '000000', 1.6, dash=False, tail=tail)
    else:
        polyline(S, pts, '9673A6', 1.4, dash=True, tail=tail)

# ---------- 块 ----------
for x, y, w, h, name, math, fc, ec, ref in D.BLOCKS:
    s = S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(ty(y, h)),
                    Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fc)
    s.line.color.rgb = RGBColor.from_string(ec)
    s.line.width = Pt(1.2)
    s.adjustments[0] = min(0.5, 0.08 / min(w, h))
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Emu(int(0.04 * EMU_IN)))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    add_runs(p0, [name], 9.5, '1A1A1A', bold=True)
    if math:
        p1 = tf.add_paragraph()
        p1.alignment = PP_ALIGN.CENTER
        add_runs(p1, math, 8, '333333')
    if ref:
        textbox(S, x + w - 0.06, y + 0.11, ref, 6.8, '777777', align='r',
                width=1.4, height=0.16)

# ---------- 圆柱 ----------
for x, y, w, h, name, ec in D.CYLINDERS:
    c = S.add_shape(MSO_SHAPE.CAN, Inches(x), Inches(ty(y, h)), Inches(w),
                    Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    c.line.color.rgb = RGBColor.from_string(ec)
    c.line.width = Pt(1.1)
    c.adjustments[0] = 0.28
    c.shadow.inherit = False
    tf = c.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, seg in enumerate(name.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, [seg], 8, '333333')

# ---------- 标签 ----------
for x, y, runs, color, fs, ha, rot in D.LABELS:
    textbox(S, x, y, runs, fs, color, align=ha, rot=rot)

# ---------- 图例 ----------
polyline(S, D.LEGEND['rt_line'], '000000', 1.6, tail=False)
lx, lyy, lt = D.LEGEND['rt_text']
textbox(S, lx, lyy, lt, 8, '000000', align='l', width=2.2)
polyline(S, D.LEGEND['sl_line'], '9673A6', 1.4, dash=True, tail=False)
lx, lyy, lt = D.LEGEND['sl_text']
textbox(S, lx, lyy, lt, 8, '000000', align='l', width=2.2)

prs.save('framework_editable.pptx')
print('saved framework_editable.pptx, shapes:', len(S._spTree))
